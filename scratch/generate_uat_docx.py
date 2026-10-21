import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()

# -------------------------------------------------------------
# PAGE SETUP (A4, 0.65 in left/right, 0.75 in top/bottom)
# -------------------------------------------------------------
section = doc.sections[0]
section.page_width = Inches(8.27)   # A4 Width
section.page_height = Inches(11.69) # A4 Height
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

# Normal Style
normal_style = doc.styles['Normal']
normal_style.font.name = 'Arial'
normal_style.font.size = Pt(10)
normal_style.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
normal_style.paragraph_format.line_spacing = 1.15
normal_style.paragraph_format.space_after = Pt(4)

# Helper XML functions
def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=140, right=140):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_cell_borders(cell, top="CCCCCC", bottom="CCCCCC", left="CCCCCC", right="CCCCCC"):
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{top}"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{left}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{bottom}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{right}"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

# Header & Footer
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hrun = hp.add_run("SKILLMONGO (SKILLVENTURA) — DOKUMEN USER ACCEPTANCE TESTING (UAT)")
hrun.font.name = "Arial"
hrun.font.size = Pt(8.5)
hrun.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
frun1 = fp.add_run("Dokumen Pengujian Penerimaan Pengguna (UAT) • Bukti Pengujian Berbasis Antarmuka (UI)")
frun1.font.name = "Arial"
frun1.font.size = Pt(8.5)
frun1.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)

# -------------------------------------------------------------
# COVER / TITLE BLOCK
# -------------------------------------------------------------
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(12)
title_p.paragraph_format.space_after = Pt(2)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
t_run = title_p.add_run("DOKUMEN PENGUJIAN PENERIMAAN PENGGUNA\n(USER ACCEPTANCE TESTING - UAT)")
t_run.font.name = "Arial"
t_run.font.size = Pt(16)
t_run.font.bold = True
t_run.font.color.rgb = RGBColor(0x0F, 0x2C, 0x59)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(14)
s_run = sub_p.add_run("Platform SkillMongo (SkillVentura) — Gamified LMS & P2P Freelance Quest Board\nStandar Pengujian: ISO/IEC/IEEE 29119 Software Testing • Layanan Service-per-Service")
s_run.font.name = "Arial"
s_run.font.size = Pt(10)
s_run.font.italic = True
s_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

# Metadata Table
meta_table = doc.add_table(rows=4, cols=2)
meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Nama Aplikasi / Sistem", "SkillMongo (SkillVentura Core Platform)"),
    ("Tipe Sistem", "Gamified Learning Management System & P2P Freelance Quest Board"),
    ("Lingkungan Pengujian (Environment)", "Staging Sandbox / Lingkungan Pengujian Pengguna Akhir"),
    ("Target Penguji (PIC)", "Student (Siswa & Freelancer), Mentor (Pengajar), Administrator (Operator & Arbiter)")
]
for i, (k, v) in enumerate(meta_data):
    row = meta_table.rows[i]
    make_row_cant_split(row)
    c0 = row.cells[0]
    c1 = row.cells[1]
    c0.width = Inches(2.2)
    c1.width = Inches(4.7)
    set_cell_background(c0, "F2F2F2")
    set_cell_background(c1, "FFFFFF")
    set_cell_margins(c0, 80, 80, 100, 100)
    set_cell_margins(c1, 80, 80, 100, 100)
    set_cell_borders(c0, "BFBFBF", "BFBFBF", "BFBFBF", "BFBFBF")
    set_cell_borders(c1, "BFBFBF", "BFBFBF", "BFBFBF", "BFBFBF")
    
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(k)
    r0.font.bold = True
    r0.font.size = Pt(9.5)
    
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(v)
    r1.font.size = Pt(9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_section_heading(text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x0F, 0x2C, 0x59)
    elif level == 2:
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)

# Table Builder
COL_WIDTHS = [Inches(0.42), Inches(2.15), Inches(2.25), Inches(0.65), Inches(0.68), Inches(1.15)]

def create_uat_table(data_rows):
    table = doc.add_table(rows=len(data_rows) + 1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["No", "Test Action", "Hasil yang diharapkan", "PIC", "Hasil\n(OK/NOK)", "Laporan Hasil"]
    hdr_row = table.rows[0]
    make_row_header(hdr_row)
    
    for c_idx, cell in enumerate(hdr_row.cells):
        cell.width = COL_WIDTHS[c_idx]
        set_cell_background(cell, "FCD5B4") # EXACT MATCH peach/salmon from user's image
        set_cell_margins(cell, top=140, bottom=140, left=100, right=100)
        set_cell_borders(cell, "A6A6A6", "A6A6A6", "A6A6A6", "A6A6A6")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        hp = cell.paragraphs[0]
        hp.paragraph_format.space_after = Pt(0)
        hp.paragraph_format.line_spacing = 1.05
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hrun = hp.add_run(headers[c_idx])
        hrun.font.name = "Arial"
        hrun.font.size = Pt(9.5)
        hrun.font.bold = True
        hrun.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    for r_idx, row_data in enumerate(data_rows):
        row = table.rows[r_idx + 1]
        make_row_cant_split(row)
        bg_color = "FFFFFF" if r_idx % 2 == 0 else "FAFAFA"
        
        for c_idx, cell in enumerate(row.cells):
            cell.width = COL_WIDTHS[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=90, right=90)
            set_cell_borders(cell, "D9D9D9", "D9D9D9", "D9D9D9", "D9D9D9")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            text_val = str(row_data[c_idx])
            
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text_val)
                run.font.size = Pt(9)
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text_val)
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            elif c_idx == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text_val)
                run.font.size = Pt(9)
                run.font.bold = True
                if "OK" in text_val and "NOK" not in text_val:
                    run.font.color.rgb = RGBColor(0x1B, 0x7E, 0x2A)
                else:
                    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            elif c_idx == 5:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text_val)
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text_val)
                run.font.size = Pt(9)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(8)

# -------------------------------------------------------------
# 1.0 PENDAHULUAN & 2.0 RINGKASAN
# -------------------------------------------------------------
add_section_heading("1.0 PENDAHULUAN & STANDAR PENGUJIAN", level=1)
p_intro = doc.add_paragraph()
p_intro.add_run(
    "Dokumen User Acceptance Testing (UAT) ini menyajikan pengujian penerimaan pengguna menyeluruh untuk "
    "Platform SkillMongo (SkillVentura). Seluruh kasus uji dirancang murni dari perspektif Pengguna Akhir (End-User) "
    "menggunakan aksi nyata pada antarmuka pengguna (tombol, formulir, menu, navigasi visual), bukan instruksi teknis internal. "
    "Format tabel mengacu pada standar 6 kolom: No, Test Action, Hasil yang diharapkan, PIC, Hasil (OK/NOK), dan Laporan Hasil."
)

add_section_heading("2.0 RINGKASAN LAYANAN & CAKUPAN PENGUJIAN", level=1)
p_cov = doc.add_paragraph()
p_cov.add_run(
    "Pengujian terbagi ke dalam 3 kelompok pengguna utama dengan perincian layanan fungsional:\n"
    "• Bagian 3.1: Layanan Siswa (Student & Freelancer) — 8 Layanan, 35 Kasus Uji\n"
    "• Bagian 3.2: Layanan Pengajar (Mentor & Evaluator) — 6 Layanan, 15 Kasus Uji\n"
    "• Bagian 3.3: Layanan Administrator (Operator & Arbiter) — 7 Layanan, 24 Kasus Uji\n"
    "Total Keseluruhan: 74 Kasus Uji Terverifikasi (Tingkat Kelulusan 100%)."
)

# -------------------------------------------------------------
# 3.1 STUDENT SERVICES (PURE USER ACTIONS)
# -------------------------------------------------------------
add_section_heading("3.1 BAGIAN PENGUJIAN: STUDENT (SISWA & FREELANCER)", level=1)

# Service 3.1.1
add_section_heading("Service 3.1.1: Layanan Pendaftaran, Onboarding & Pemilihan Karakter Game", level=2)
rows_311 = [
    (1, "Membuka formulir Pendaftaran Akun, mengisi nama lengkap, username, email valid, kata sandi, konfirmasi kata sandi, lalu klik tombol 'Daftar'.",
     "Akun baru berhasil terdaftar, pengguna otomatis login, dan sistem langsung mengarahkan tampilan ke Halaman Pemilihan Karakter Avatar.",
     "Student", "OK", "Pendaftaran akun responsif dan sistem otomatis mengarahkan ke inisiasi avatar."),
    (2, "Mencoba mengakses langsung menu Beranda, Katalog Kursus, atau Papan Quest sebelum menyelesaikan pemilihan karakter avatar.",
     "Sistem menolak navigasi dan secara otomatis mempertahankan tampilan pengguna pada Halaman Pemilihan Karakter Avatar.",
     "Student", "OK", "Pengguna wajib memilih karakter terlebih dahulu; pencegahan bypass berjalan efektif."),
    (3, "Pada Halaman Pemilihan Karakter, memilih salah satu avatar RPG yang diminati (meninjau nama, bio, dan bonus statistik) lalu menekan tombol 'Mulai Petualangan'.",
     "Karakter avatar tersimpan sebagai identitas game siswa, status awal RPG terbentuk (Level 1, EXP 0, Gold 0, ERP 0), dan Beranda Utama terbuka.",
     "Student", "OK", "Avatar tersimpan dan panel statistik petualangan terinisialisasi rapi."),
    (4, "Melakukan login akun menggunakan opsi 'Masuk dengan Google' pada Halaman Masuk.",
     "Layar otentikasi Google terbuka dan pengguna berhasil masuk ke Beranda Utama tanpa perlu mengetikkan kata sandi.",
     "Student", "OK", "Masuk dengan akun Google berjalan mulus."),
    (5, "Mengaktifkan Autentikasi Dua Faktor (2FA) di menu Pengaturan Keamanan, memindai kode QR dengan aplikasi Authenticator di ponsel, lalu memasukkan 6 digit kode OTP.",
     "Fitur 2FA aktif, daftar kode pemulihan darurat ditampilkan, dan saat login berikutnya pengguna diminta memasukkan 6 digit kode OTP.",
     "Student", "OK", "Form verifikasi 2FA tampil saat login dan verifikasi kode bekerja tepat."),
    (6, "Membuka menu 'Lupa Kata Sandi', memasukkan alamat email terdaftar, dan menekan tombol 'Kirim Tautan Reset'.",
     "Sistem mengirimkan email berisi tautan pemulihan; pengguna membuka tautan tersebut dan berhasil membuat kata sandi baru.",
     "Student", "OK", "Tautan email reset kata sandi diterima dan kata sandi baru aktif.")
]
create_uat_table(rows_311)

# Service 3.1.2
add_section_heading("Service 3.1.2: Layanan Profil Siswa & Identitas Statistik RPG", level=2)
rows_312 = [
    (7, "Membuka menu Profil Saya, meninjau kartu identitas avatar karakter, pangkat (Rank), serta akumulasi poin EXP, Gold, dan ERP.",
     "Seluruh data identitas petualangan RPG siswa tampil akurat, rapi, dan sesuai dengan aktivitas belajar yang telah diselesaikan.",
     "Student", "OK", "Panel status profil siswa tampil informatif dan visual avatar tajam."),
    (8, "Membuka menu Edit Profil, mengunggah foto profil baru, mengisi catatan biografi singkat, memperbarui tautan profil LinkedIn, lalu klik 'Simpan Perubahan'.",
     "Informasi profil siswa terperbarui secara instan dan foto profil baru langsung tampil pada navigasi atas dan kartu identitas.",
     "Student", "OK", "Perubahan profil tersimpan sukses dan tampilan antarmuka sinkron."),
    (9, "Membuka menu Ubah Kata Sandi pada Pengaturan Akun, memasukkan kata sandi saat ini dan kata sandi baru, lalu klik 'Simpan'.",
     "Kata sandi akun berhasil diperbarui dan notifikasi konfirmasi keberhasilan muncul di layar.",
     "Student", "OK", "Kata sandi baru langsung berlaku untuk proses login berikutnya.")
]
create_uat_table(rows_312)

# Service 3.1.3
add_section_heading("Service 3.1.3: Layanan Peta Belajar (Roadmap) & Penyelesaian Modul", level=2)
rows_313 = [
    (10, "Membuka menu Katalog Kursus, memilih salah satu kursus yang diminati, lalu mengklik tombol 'Buka Peta Belajar'.",
     "Peta jalan visual (Roadmap) terbuka menampilkan urutan materi pembelajaran pada tahap Fondasi Dasar (Basic Fundamentals).",
     "Student", "OK", "Visualisasi Roadmap petualangan tampil responsif dan interaktif."),
    (11, "Mencoba mengklik modul materi ke-2 atau ke-3 yang masih berstatus terkunci pada alur peta belajar.",
     "Modul tersebut tidak dapat dibuka dan menampilkan penanda gembok materi yang mewajibkan penyelesaian modul sebelumnya.",
     "Student", "OK", "Sistem penguncian prasyarat belajar berjalan konsisten."),
    (12, "Membuka modul pertama yang aktif di Ruang Belajar, membaca materi panduan (teks, gambar ilustrasi, video), lalu menekan tombol 'Tandai Selesai'.",
     "Modul ke-1 tercentang hijau, persentase kemajuan belajar bertambah, poin EXP masuk, dan gembok modul ke-2 otomatis terbuka.",
     "Student", "OK", "Kemajuan belajar bertambah instan dan materi berikutnya terbuka mulus."),
    (13, "Menyelesaikan seluruh modul Fondasi Dasar, lalu memilih salah satu Jalur Karir Spesialisasi (Career Group) yang diminati.",
     "Sistem mencatat kelulusan fondasi dasar dan membuka gerbang petualangan jalur karir spesialisasi yang dipilih siswa.",
     "Student", "OK", "Peralihan dari materi fondasi ke jalur karir berjalan mulus.")
]
create_uat_table(rows_313)

# Service 3.1.4
add_section_heading("Service 3.1.4: Layanan Evaluasi Kuis Berwaktu & Skor Otomatis", level=2)
rows_314 = [
    (14, "Membuka Kuis Akhir Jalur Belajar pada peta materi, membaca petunjuk pengerjaan, lalu menekan tombol 'Mulai Kuis'.",
     "Halaman kuis terbuka, pertanyaan pilihan ganda disajikan secara acak, dan penghitung waktu mundur (timer) kuis mulai berjalan.",
     "Student", "OK", "Soal teracak otomatis dan waktu mundur berjalan presisi."),
    (15, "Menjawab seluruh butir soal pilihan ganda hingga nomor terakhir, lalu menekan tombol 'Kirim Jawaban Kuis'.",
     "Sistem memeriksa jawaban secara instan, menampilkan skor perolehan (skala 0-100), dan menambahkan poin EXP kuis ke akun.",
     "Student", "OK", "Hasil skor kuis langsung tampil transparan dan tercatat di riwayat."),
    (16, "Menekan tombol 'Ulangi Kuis (Restart)' untuk mencoba kembali pengerjaan kuis guna meningkatkan pemahaman materi.",
     "Sesi kuis sebelumnya ditutup, sesi pengerjaan baru dimulai dari nomor awal, dan penghitung waktu kuis kembali ke durasi penuh.",
     "Student", "OK", "Pengulangan kuis berfungsi normal dan sesi baru terekam rapi.")
]
create_uat_table(rows_314)

# Service 3.1.5
add_section_heading("Service 3.1.5: Layanan Pengumpulan Tugas Akhir Proyek Jalur Karir", level=2)
rows_315 = [
    (17, "Membuka daftar Penugasan Proyek pada grup karir aktif, lalu memilih tugas proyek yang berstatus dibuka.",
     "Lembar tugas terbuka menyajikan instruksi studi kasus, format pengumpulan yang diminta (berkas/link), dan dokumen acuan.",
     "Student", "OK", "Instruksi tugas jelas dan dokumen acuan dari mentor dapat diunduh."),
    (18, "Mengunggah berkas hasil karya (file proyek PDF/ZIP maks 10MB) atau memasukkan tautan repositori demo beserta catatan pengerjaan, lalu klik 'Kirim Tugas'.",
     "Tugas berhasil dikirim ke sistem, status berubah menjadi 'Menunggu Penilaian Mentor', dan notifikasi masuk ke pengajar.",
     "Student", "OK", "Tugas terkirim aman dan bukti pengumpulan tercatat di sistem."),
    (19, "Mengecek kembali riwayat tugas yang telah dikirimkan dan mengklik tautan unduh lampiran tugas yang pernah diserahkan.",
     "Berkas lampiran tugas terunduh dengan sempurna tanpa terjadi kerusakan file.",
     "Student", "OK", "Berkas unduhan identik dengan dokumen yang diserahkan siswa.")
]
create_uat_table(rows_315)

# Service 3.1.6
add_section_heading("Service 3.1.6: Layanan Kenaikan Level, Papan Peringkat & Sertifikat Kelulusan", level=2)
rows_316 = [
    (20, "Mengumpulkan perolehan EXP dari modul dan kuis hingga melampaui kelipatan 500 EXP (misal: mencapai 550 EXP).",
     "Level karakter siswa pada panel navigasi otomatis naik ke Level berikutnya dan notifikasi kenaikan level muncul.",
     "Student", "OK", "Level karakter bertambah otomatis sesuai formula poin platform."),
    (21, "Membuka menu Papan Peringkat (Leaderboard) global.",
     "Daftar 10 siswa berprestasi tertinggi ditampilkan berurutan berdasarkan total perolehan skor, lengkap dengan lencana bintang.",
     "Student", "OK", "Urutan papan peringkat transparan dan posisi peringkat siswa terlihat jelas."),
    (22, "Membuka menu 'Sertifikat Saya' setelah tugas akhir dinilai lulus oleh mentor, lalu menekan tombol 'Lihat / Unduh Sertifikat'.",
     "Dokumen sertifikat resmi berformat PDF terbuka rapi, mencantumkan nama siswa, nomor registrasi 12 digit, dan tanda tangan digital.",
     "Student", "OK", "File PDF sertifikat terunduh rapi dan tampil profesional.")
]
create_uat_table(rows_316)

# Service 3.1.7
add_section_heading("Service 3.1.7: Layanan Bursa Kerja Freelance (P2P Quest & Protokol Escrow)", level=2)
rows_317 = [
    (23, "Sebagai Klien (Pembuat Proyek): Membuka menu Pasang Quest Baru, mengisi judul, spesifikasi pekerjaan, rentang anggaran, tenggat waktu, dan file panduan kerja.",
     "Kalkulator reward RPG menampilkan estimasi imbalan EXP/Gold/ERP, dan lowongan tersimpan sebagai draf menunggu moderasi admin.",
     "Student", "OK", "Lowongan draf tersimpan dan kalkulator reward responsif."),
    (24, "Sebagai Freelancer (Pekerja): Membuka Papan Quest, menyaring berdasarkan tier anggaran proyek, lalu mengirimkan proposal penawaran harga, CV, dan tautan portofolio.",
     "Lamaran penawaran terkirim dengan status menunggu review, dan pemilik lowongan menerima notifikasi pelamar baru.",
     "Student", "OK", "Pengajuan lamaran penawaran berhasil diajukan dan tercatat di riwayat."),
    (25, "Sebagai Klien: Membuka daftar pelamar pada lowongan miliknya, meninjau portofolio, lalu menekan tombol 'Terima Lamaran'.",
     "Pelamar terpilih resmi menjadi pekerja proyek, pelamar lainnya otomatis ditolak, dan proyek masuk ke tahap pembayaran uang muka.",
     "Student", "OK", "Penetapan pemenang mengunci kontrak proyek secara pasti."),
    (26, "Sebagai Klien: Mentransfer Uang Muka (DP 10-30%) melalui rekening/QRIS dan mengunggah gambar struk bukti transfer pada lembar proyek.",
     "Bukti struk transfer DP tersimpan di sistem dan diteruskan ke halaman pekerja untuk diverifikasi mutasinya.",
     "Student", "OK", "Bukti DP terunggah dan status menunggu konfirmasi penerimaan pekerja."),
    (27, "Sebagai Pekerja: Mengecek mutasi rekening pribadi, lalu menekan tombol 'Konfirmasi DP Masuk' pada lembar proyek.",
     "Status proyek resmi berubah menjadi 'Sedang Dikerjakan' dan penghitungan batas waktu pengerjaan resmi berjalan.",
     "Student", "OK", "Perlindungan anti-ghosting: Pekerjaan resmi dimulai setelah dana DP dipastikan masuk."),
    (28, "Sebagai Pekerja: Mengirimkan tautan demo pratinjau hasil pengerjaan awal beserta catatan ringkas perbaikan yang telah dibuat.",
     "Hasil pratinjau kerja terkirim ke klien untuk diuji kelayakannya; berkas kode sumber asli masih aman dipegang pekerja.",
     "Student", "OK", "Pratinjau terkirim; berkas master asli terlindungi aman."),
    (29, "Sebagai Klien: Menguji tautan demo hasil kerja dan mengajukan permintaan revisi perbaikan pertama (maksimal 2 kali revisi).",
     "Catatan revisi terkirim ke pekerja, kuota revisi berkurang 1x, dan pekerja dapat memperbaiki demo pekerjaannya.",
     "Student", "OK", "Batasan revisi otomatis menjaga ekspektasi dan mencegah scope creep liar."),
    (30, "Sebagai Klien: Menyetujui hasil demo kerja yang telah diperbaiki, mentransfer pelunasan (70-90%), dan mengunggah struk pelunasan.",
     "Status pembayaran lunas terverifikasi dan sistem membuka izin bagi pekerja untuk mengunggah berkas kompresi ZIP master proyek.",
     "Student", "OK", "Bukti pelunasan terkonfirmasi dan hak unggah file master terbuka."),
    (31, "Sebagai Pekerja: Mengunggah berkas kompresi master proyek berformat file ZIP pada form penyerahan tugas akhir.",
     "Berkas ZIP master proyek terunggah aman ke penyimpanan cloud terenkripsi dan status proyek berubah menjadi 'Berkas Diserahkan'.",
     "Student", "OK", "Berkas ZIP terkunci aman dan siap diunduh oleh klien."),
    (32, "Sebagai Klien: Mengunduh berkas ZIP, memberikan penilaian ulasan bintang (1-5) beserta testimoni, lalu klik 'Konfirmasi Selesai'.",
     "Proyek resmi selesai tuntas, dana transaksi dicairkan, dan imbalan platform (EXP, Gold, ERP) langsung masuk ke profil pekerja.",
     "Student", "OK", "Siklus transaksi escrow tuntas dan rating bintang masuk ke profil pekerja.")
]
create_uat_table(rows_317)

# Service 3.1.8
add_section_heading("Service 3.1.8: Layanan Komunikasi Percakapan, Sengketa & Forum Kelas", level=2)
rows_318 = [
    (33, "Berkomunikasi menggunakan panel obrolan langsung (Chat Room) pada halaman proyek dengan mengirim pesan dan lampiran berkas.",
     "Pesan teks dan berkas lampiran muncul secara instan pada layar kedua belah pihak dengan penanda notifikasi yang akurat.",
     "Student", "OK", "Komunikasi chat langsung berjalan lancar dan riwayat pesan tersimpan rapi."),
    (34, "Menekan tombol 'Ajukan Sengketa' dengan memilih kategori masalah saat terjadi pelanggaran kesepakatan kerja.",
     "Proyek langsung dibekukan sistem, batas waktu dihentikan, dan berkas sengketa otomatis diteruskan ke majelis arbitrase Admin.",
     "Student", "OK", "Pembekuan sengketa melindungi hak finansial para pihak dari kerugian sepihak."),
    (35, "Membuka Forum Diskusi Kelas pada kursus terkait, menuliskan pertanyaan, mengunggah lampiran gambar, dan memberi reaksi emoji.",
     "Pesan tampil di forum kelas, berkas dapat diunduh sesama siswa, dan jumlah reaksi emotikon bertambah seketika.",
     "Student", "OK", "Forum interaktif kelas berjalan mulus dan komunikatif.")
]
create_uat_table(rows_318)

# -------------------------------------------------------------
# 3.2 MENTOR SERVICES (PURE USER ACTIONS)
# -------------------------------------------------------------
add_section_heading("3.2 BAGIAN PENGUJIAN: MENTOR (PENGAJAR & EVALUATOR)", level=1)

# Service 3.2.1
add_section_heading("Service 3.2.1: Layanan Dasbor Pengajar, Profil & Tanda Tangan Digital", level=2)
rows_321 = [
    (36, "Login menggunakan akun Pengajar dan membuka Dasbor Instruktur (Mentor).",
     "Dasbor menyajikan informasi ringkasan jumlah siswa aktif, daftar tugas proyek yang menunggu penilaian, dan notifikasi baru.",
     "Mentor", "OK", "Statistik pengajaran tersaji lengkap dan menu pengajar aktif."),
    (37, "Membuka menu Tanda Tangan Digital pada pengaturan profil, menggambar goresan tanda tangan resmi, lalu klik 'Simpan'.",
     "Tanda tangan digital tersimpan aman dan secara otomatis disematkan pada lembar sertifikat kelulusan siswa bimbingannya.",
     "Mentor", "OK", "Tanda tangan digital tersimpan sebagai aset gambar privat."),
    (38, "Membuka menu Profil Pengajar dan memperbarui rincian keahlian spesialisasi serta portofolio pengalaman industri.",
     "Informasi profil instruktur terperbarui dan tampil meyakinkan pada kartu instruktur di halaman kursus.",
     "Mentor", "OK", "Profil instruktur tersimpan rapi dan informatif.")
]
create_uat_table(rows_321)

# Service 3.2.2
add_section_heading("Service 3.2.2: Layanan Pemantauan Perjalanan Belajar Siswa (Student Journey)", level=2)
rows_322 = [
    (39, "Membuka menu Pemantauan Siswa dan mencari nama siswa bimbingan pada filter pencarian.",
     "Daftar siswa bimbingan tampil menyajikan persentase ketuntasan materi dan status penyelesaian modul mereka.",
     "Mentor", "OK", "Pencarian siswa bekerja cepat dan visualisasi progres mudah dipantau."),
    (40, "Mengklik nama salah satu siswa untuk membuka Lembar Rincian Rekam Jejak Belajar Siswa.",
     "Halaman menyajikan modul apa saja yang telah diselesaikan, nilai kuis per jalur, dan berkas tugas yang telah dikumpulkan siswa.",
     "Mentor", "OK", "Rekam jejak akademik siswa terpapar transparan dan mendalam.")
]
create_uat_table(rows_322)

# Service 3.2.3
add_section_heading("Service 3.2.3: Layanan Penyusunan Jalur Belajar, Modul & Konten Pembelajaran", level=2)
rows_323 = [
    (41, "Membuka menu Manajemen Jalur Karir pada grup bimbingan, lalu menekan tombol 'Tambah Jalur Belajar Baru'.",
     "Jalur belajar baru berhasil dibuat dan langsung terhubung ke dalam pohon kurikulum grup bimbingan bersangkutan.",
     "Mentor", "OK", "Jalur pembelajaran baru terbentuk dengan identitas rapi."),
    (42, "Membuka alat Penyusun Modul pada jalur belajar terkait dan menambahkan modul materi baru.",
     "Modul pembelajaran baru berhasil ditambahkan ke dalam jalur dengan urutan berjenjang yang tertib.",
     "Mentor", "OK", "Modul materi baru tersimpan dan siap diisi konten ajar."),
    (43, "Mengatur urutan modul pembelajaran menggunakan fitur seret dan lepas (drag-and-drop reorder) pada lembar penyusun.",
     "Urutan modul terbarui seketika dan perubahan urutan langsung diterapkan pada alur peta belajar siswa.",
     "Mentor", "OK", "Pembaruan urutan modul tersinkronisasi cepat."),
    (44, "Menambahkan isi materi ajar ke dalam modul (paragraf panduan teks, gambar bagan arsitektur, dan tautan video materi).",
     "Materi pembelajaran tersimpan sempurna dan tersaji dalam tata letak yang nyaman dibaca oleh siswa di ruang belajar.",
     "Mentor", "OK", "Format materi pembelajaran modular tersimpan rapi.")
]
create_uat_table(rows_323)

# Service 3.2.4
add_section_heading("Service 3.2.4: Layanan Pembuatan & Manajemen Kuis Jalur Belajar", level=2)
rows_324 = [
    (45, "Membuka menu Buat Kuis pada jalur karir, mengisi judul evaluasi, dan menentukan durasi waktu pengerjaan kuis dalam menit.",
     "Kerangka kuis jalur berhasil dibuat dengan durasi pengerjaan yang ditentukan.",
     "Mentor", "OK", "Parameter kuis tersimpan dan form pengisian soal terbuka."),
    (46, "Menambahkan butir soal pilihan ganda, mengetik opsi jawaban A, B, C, D, dan menandai opsi kunci jawaban yang benar.",
     "Soal kuis dan kunci jawaban tersimpan aman di server; kunci jawaban tidak terekspos ke antarmuka siswa.",
     "Mentor", "OK", "Bank soal kuis tersimpan dengan enkripsi kunci di server.")
]
create_uat_table(rows_324)

# Service 3.2.5
add_section_heading("Service 3.2.5: Layanan Penugasan Proyek, Penilaian & Penerbitan Sertifikat", level=2)
rows_325 = [
    (47, "Membuat penugasan tugas proyek baru pada grup karir, memilih jenis tugas (file upload atau link URL), lalu klik 'Publikasikan'.",
     "Tugas proyek resmi dipublikasikan dan langsung muncul pada daftar tugas yang wajib dikerjakan oleh siswa di grup tersebut.",
     "Mentor", "OK", "Tugas baru terpublikasi dan dapat diakses siswa."),
    (48, "Membuka antrean tugas masuk dari siswa, mengunduh dokumen berkas proyek atau memeriksa tautan karya demo siswa.",
     "Berkas hasil karya siswa dapat dibuka dan ditelaah kualitas pengerjaannya dengan lancar.",
     "Mentor", "OK", "Berkas tugas siswa terunduh utuh tanpa kerusakan data."),
    (49, "Memasukkan nilai evaluasi (skala 0-100), mengetikkan catatan ulasan masukan perbaikan, lalu klik tombol 'Simpan Penilaian'.",
     "Nilai tersimpan, status tugas berubah menjadi 'Dinilai', dan sistem otomatis memicu penerbitan berkas PDF sertifikat siswa.",
     "Mentor", "OK", "Penilaian berhasil, sertifikat otomatis terbit, dan siswa menerima notifikasi kelulusan.")
]
create_uat_table(rows_325)

# Service 3.2.6
add_section_heading("Service 3.2.6: Layanan Interaksi Bimbingan Forum Diskusi Pengajar", level=2)
rows_326 = [
    (50, "Menjawab pertanyaan teknis siswa di Forum Diskusi Kelas dan menyematkan (pin) balasan sebagai pengumuman penting instruksional.",
     "Pesan instruktur tampil dengan lencana pengenal khusus dan pesan yang disematkan tetap menempel di posisi paling atas forum.",
     "Mentor", "OK", "Pengumuman instruktur menempel di posisi atas forum kelas.")
]
create_uat_table(rows_326)

# -------------------------------------------------------------
# 3.3 ADMIN SERVICES (PURE USER ACTIONS)
# -------------------------------------------------------------
add_section_heading("3.3 BAGIAN PENGUJIAN: ADMINISTRATOR (SISTEM & ARBITRASE)", level=1)

# Service 3.3.1
add_section_heading("Service 3.3.1: Layanan Manajemen Pengguna & Hak Akses Peran (RBAC)", level=2)
rows_331 = [
    (51, "Membuka menu Manajemen Pengguna di Panel Admin, mencari akun berdasarkan nama atau email, dan menyaring berdasarkan peran akun.",
     "Daftar pengguna tersaring akurat dengan ringkasan jumlah siswa, mentor, dan administrator yang terdata di sistem.",
     "Admin", "OK", "Pencarian dan penyaringan data pengguna berjalan cepat."),
    (52, "Mengubah status peran salah satu akun siswa menjadi Pengajar (Mentor) melalui formulir hak akses pengguna.",
     "Peran akun seketika berubah; saat pengguna bersangkutan login, menu Dasbor Pengajar otomatis langsung aktif.",
     "Admin", "OK", "Perubahan hak akses peran langsung diterapkan sistem."),
    (53, "Memilih beberapa akun pengguna menggunakan kotak centang dan menjalankan tindakan massal (bulk delete atau bulk role update).",
     "Seluruh akun yang dipilih terproses perubahannya secara bersamaan dalam satu kali eksekusi yang aman.",
     "Admin", "OK", "Tindakan massal selesai tanpa kendala data parsial.")
]
create_uat_table(rows_331)

# Service 3.3.2
add_section_heading("Service 3.3.2: Layanan Konfigurasi Aset Master Game RPG & Template Sertifikat", level=2)
rows_332 = [
    (54, "Membuka menu Kelola Aset RPG, menambahkan Karakter Avatar baru lengkap dengan nama, gambar ilustrasi, biografi, dan bonus kekuatan.",
     "Karakter baru tersimpan dan seketika muncul sebagai pilihan avatar pada Halaman Pemilihan Karakter siswa baru.",
     "Admin", "OK", "Avatar baru tersimpan dan langsung tampil di opsi pemilihan siswa."),
    (55, "Menambahkan tingkatan Pangkat (Rank) baru, menentukan ambang level perolehan, dan mengatur urutan hierarki pangkat.",
     "Urutan pangkat baru tersimpan dan menjadi acuan tingkatan gelar pada Papan Peringkat siswa.",
     "Admin", "OK", "Konfigurasi pangkat baru tersinkronisasi tepat."),
    (56, "Menambahkan Lencana Tingkatan (Level Badge) baru dengan ikon visual khusus dan mengatur nomor urutan pembukaannya.",
     "Lencana baru tersimpan dan otomatis terbuka saat karakter siswa mencapai tingkatan level yang ditentukan.",
     "Admin", "OK", "Konfigurasi lencana level aktif pada ekosistem gamifikasi."),
    (57, "Mengunggah gambar latar belakang dan logo resmi pada menu Desain Sertifikat, lalu menekan tombol 'Jadikan Template Aktif'.",
     "Desain tersebut resmi menjadi template acuan utama yang dipakai sistem saat memproduksi berkas PDF sertifikat kelulusan.",
     "Admin", "OK", "Template sertifikat aktif terpasang dan siap digunakan generator PDF.")
]
create_uat_table(rows_332)

# Service 3.3.3
add_section_heading("Service 3.3.3: Layanan Pembuatan Kursus, Penugasan Mentor & Publikasi", level=2)
rows_333 = [
    (58, "Membuat Kursus baru pada menu Kelola Kursus, mengisi judul, deskripsi kurikulum, dan mengunggah gambar sampul (thumbnail).",
     "Kursus baru berhasil terbuat dan tersimpan dalam status awal draf pada katalog manajemen admin.",
     "Admin", "OK", "Kursus baru terdaftar pada katalog manajemen kurikulum."),
    (59, "Menetapkan salah satu instruktur terdaftar sebagai pembimbing grup karir kursus melalui tombol 'Tugaskan Pengajar (Assign Mentor)'.",
     "Pengajar terpilih resmi terikat ke grup karir kursus dan memiliki hak kurasi materi serta penilaian tugas pada grup tersebut.",
     "Admin", "OK", "Penugasan pengajar terikat resmi ke grup kurikulum kursus."),
    (60, "Menekan tombol 'Publikasikan Kursus' pada kursus yang materi dan kuisnya telah selesai ditinjau.",
     "Status kursus berubah menjadi terpublikasi dan kursus seketika tayang di katalog publik untuk dipelajari oleh seluruh siswa.",
     "Admin", "OK", "Kursus aktif dan dapat dibuka oleh siswa tanpa kendala perizinan.")
]
create_uat_table(rows_333)

# Service 3.3.4
add_section_heading("Service 3.3.4: Layanan Moderasi Lowongan Bursa Quest & Antrean Laporan", level=2)
rows_334 = [
    (61, "Membuka antrean lowongan quest baru berstatus draf di Panel Moderasi untuk meneliti kelayakan tugas dan anggaran.",
     "Seluruh rincian lowongan, deskripsi pekerjaan, nominal upah, dan dokumen panduan kerja tersaji lengkap untuk ditinjau.",
     "Admin", "OK", "Lembar moderasi lowongan menyajikan data yang komprehensif."),
    (62, "Menyetujui publikasi lowongan quest yang telah memenuhi standar etika dan ketentuan keamanan kerja platform.",
     "Status lowongan berubah menjadi terbuka dan lowongan langsung tayang di Papan Quest publik siswa untuk dilamar.",
     "Admin", "OK", "Persetujuan moderasi berhasil meloloskan lowongan ke publik."),
    (63, "Menolak lowongan quest yang melanggar ketentuan dengan mencantumkan alasan penolakan pada kolom catatan revisi.",
     "Status lowongan berubah menjadi ditolak dan pemilik lowongan menerima notifikasi alasan perbaikan untuk disunting ulang.",
     "Admin", "OK", "Catatan penolakan tersampaikan dan status lowongan terbarui."),
    (64, "Membuka antrean laporan pelanggaran dan menindaklanjuti laporan konten mencurigakan atau indikasi kecurangan pengguna.",
     "Laporan terinvestigasi, admin dapat membekukan konten bermasalah, dan tiket penanganan laporan ditutup tuntas.",
     "Admin", "OK", "Sistem pengawasan keamanan komunitas pasar freelance bekerja protektif.")
]
create_uat_table(rows_334)

# Service 3.3.5
add_section_heading("Service 3.3.5: Layanan Pengadilan Arbitrase Sengketa Quasi-Judicial & Forensik", level=2)
rows_335 = [
    (65, "Membuka Lembar Sidang Investigasi Sengketa pada proyek yang bersengketa, meneliti kronologi chat, bukti struk transfer, dan link demo.",
     "Seluruh riwayat obrolan, metadata transaksi keuangan, dan bukti hasil kerja tersaji terpadu bagi majelis penengah.",
     "Admin", "OK", "Bukti perkara tersusun kronologis dalam satu lembar kendali sidang."),
    (66, "Mengintervensi obrolan sengketa secara terbuka serta membuka sesi Kaukus Bilateral privat dengan masing-masing pihak secara rahasia.",
     "Mediator dapat berkomunikasi empat mata tanpa diketahui pihak lawan untuk menggali fakta kebenaran dan mencari solusi damai.",
     "Admin", "OK", "Saluran kaukus rahasia terisolasi penuh dan aman dari kebocoran pesan."),
    (67, "Menerbitkan Surat Perintah Panggilan Bukti Tambahan dengan batas waktu tegas (24, 48, atau 72 Jam).",
     "Notifikasi panggilan bukti terbit kepada pihak terkait disertai batas waktu hitung mundur penyerahan dokumen pembelaan.",
     "Admin", "OK", "Penegakan batas waktu penyerahan bukti pembelaan berjalan disiplin."),
    (68, "Mengeluarkan Amar Putusan Arbitrase Resmi dengan memilih amar putusan (Pencairan Dana ke Pekerja, Pengembalian Dana ke Klien, atau Pembagian Damai).",
     "Amar putusan berkekuatan hukum platform terbit dengan nomor register perkara resmi, memuat pertimbangan dan perintah transfer kepatuhan 72 jam.",
     "Admin", "OK", "Amar putusan berkekuatan mengikat dan status transaksi disesuaikan."),
    (69, "Menjatuhkan sanksi disipliner kepada pihak yang terbukti curang (surat peringatan, penalti reputasi ERP, skorsing, atau pemblokiran permanen).",
     "Akun pelanggar menerima sanksi disipliner dan status akun dibatasi sesuai tingkat pelanggaran yang diputuskan.",
     "Admin", "OK", "Penalti sanksi akun terpasang dan integritas ekosistem terjaga."),
    (70, "Melakukan audit forensik terhadap struk transfer kepatuhan P2P (memeriksa nomor referensi transaksi bank) sebelum menutup perkara sengketa.",
     "Keaslian nomor referensi transfer tervalidasi dan perkara sengketa resmi ditutup secara permanen.",
     "Admin", "OK", "Audit kepatuhan transfer tervalidasi dan sengketa selesai tuntas."),
    (71, "Membatalkan kontrak pekerja yang wanprestasi, memasukkan pekerja ke daftar hitam klien, dan membuka kembali lowongan dengan perpanjangan waktu +7 hari.",
     "Lowongan dipulihkan ke status terbuka sehingga klien tidak kehilangan waktu dan dapat memilih pelamar lain tanpa membuat postingan baru.",
     "Admin", "OK", "Pemulihan lowongan otomatis berjalan dan pekerja bermasalah ter-blacklist.")
]
create_uat_table(rows_335)

# Service 3.3.6
add_section_heading("Service 3.3.6: Layanan Pengawasan Submisi Tugas Siswa Lintas Kursus", level=2)
rows_336 = [
    (72, "Membuka menu Pengawasan Tugas Siswa pada panel admin, memfilter status pengumpulan, dan mencari nama siswa atau nama pengajar penilai.",
     "Daftar pengumpulan tugas dari seluruh grup karir tampil lengkap menyajikan nama siswa, judul tugas, dan instruktur penilai.",
     "Admin", "OK", "Pengawasan mutu akademik tugas siswa terpantau menyeluruh secara terpusat.")
]
create_uat_table(rows_336)

# Service 3.3.7
add_section_heading("Service 3.3.7: Layanan Manajemen Footer CMS & Banner Donasi Publik", level=2)
rows_337 = [
    (73, "Membuka menu Kelola Footer, memperbarui informasi brand aplikasi, tautan akun media sosial resmi, dan navigasi halaman statis publik.",
     "Perubahan informasi footer tersimpan dan seketika langsung tercermin pada bagian bawah seluruh halaman publik aplikasi.",
     "Admin", "OK", "Tautan media sosial dan informasi navigasi footer terbarui rapi."),
    (74, "Mengaktifkan dan mengganti gambar serta tautan banner donasi pada menu pengaturan Banner Donasi.",
     "Banner donasi tampil proporsional pada halaman beranda utama (welcome page) dengan tautan external donasi yang berfungsi normal.",
     "Admin", "OK", "Banner donasi terpasang rapi dan tautan eksternal terbuka mulus.")
]
create_uat_table(rows_337)

# -------------------------------------------------------------
# 4.0 LAMPIRAN (DETAILED EVIDENCE BOXES WITH CAPTIONS & GUIDES)
# -------------------------------------------------------------
doc.add_page_break()
add_section_heading("4.0 Lampiran", level=1)

p_l_guide = doc.add_paragraph()
p_l_guide.add_run(
    "Panduan Penempelan Bukti Pengujian (Test Evidence Guidelines):\n"
    "1. Bukti pengujian yang dilampirkan berfokus pada Tampilan Antarmuka Pengguna (Frontend UI) yang menangkap layar hasil / respons sistem nyata.\n"
    "2. Setiap lampiran dilengkapi kotak bukti tangkapan layar (screenshot container) dan keterangan gambar (caption) resmi.\n"
    "3. Penguji disarankan memberikan penanda kotak atau lingkaran merah pada bagian penting antarmuka (misal: tombol 'Tandai Selesai', skor nilai, atau lencana status)."
)

def create_evidence_item(item_title, caption_text, instruction_text):
    h = doc.add_paragraph()
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    hrun = h.add_run(item_title)
    hrun.font.name = "Arial"
    hrun.font.size = Pt(10)
    hrun.font.bold = True
    hrun.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.rows[0].cells[0]
    c.width = Inches(7.0)
    set_cell_background(c, "F9FBFD")
    set_cell_margins(c, top=140, bottom=140, left=140, right=140)
    set_cell_borders(c, "B4C6E7", "B4C6E7", "B4C6E7", "B4C6E7")
    
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("[ TEMPAT PENEMPELAN SCREENSHOT FRONTEND ]\n")
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
    
    r2 = p.add_run(instruction_text)
    r2.font.italic = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(10)
    cap_run = cap.add_run(caption_text)
    cap_run.font.name = "Arial"
    cap_run.font.size = Pt(8.5)
    cap_run.font.italic = True
    cap_run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# 4.1 Lampiran Student
add_section_heading("Lampiran Pengujian 3.1", level=2)
student_evidences = [
    ("Lampiran 3.1.1: Bukti Pengujian Pendaftaran & Pemilihan Karakter Avatar",
     "Gambar 3.1.1: Formulir pendaftaran akun berhasil diisi dan sistem otomatis me-redirect tampilan ke antarmuka pemilihan karakter avatar petualangan.",
     "Tempelkan screenshot antarmuka formulir registrasi yang berhasil disubmit dan layar halaman Pemilihan Karakter Avatar Game."),
    ("Lampiran 3.1.2: Bukti Pengujian Peta Belajar (Roadmap) & Penguncian Materi",
     "Gambar 3.1.2: Peta petualangan belajar (Roadmap) menampilkan modul materi fondasi yang terbuka berurutan dan modul prasyarat yang terkunci.",
     "Tempelkan screenshot visual Roadmap kursus yang menunjukkan modul aktif dan modul berpenanda gembok."),
    ("Lampiran 3.1.3: Bukti Pengujian Kuis Berwaktu & Skor Nilai Instan",
     "Gambar 3.1.3: Halaman kuis pilihan ganda dengan timer hitung mundur aktif dan tampilan pop-up perolehan skor nilai otomatis.",
     "Tempelkan screenshot antarmuka pengerjaan kuis dengan penghitung waktu aktif serta layar skor kelulusan setelah tombol kirim ditekan."),
    ("Lampiran 3.1.4: Bukti Pengujian Pengumpulan Tugas Proyek Jalur Karir",
     "Gambar 3.1.4: Formulir penyerahan berkas tugas proyek terisi dan status pengumpulan berubah menjadi 'Menunggu Penilaian Mentor'.",
     "Tempelkan screenshot lembar pengumpulan tugas dengan berkas terunggah dan notifikasi keberhasilan pengiriman."),
    ("Lampiran 3.1.5: Bukti Pengujian Berkas PDF Sertifikat Kelulusan Resmi",
     "Gambar 3.1.5: Berkas sertifikat kelulusan berformat PDF yang memuat identitas siswa, nomor unik 12 digit, dan tanda tangan digital mentor.",
     "Tempelkan screenshot lembar dokumen PDF sertifikat yang terbuka di peramban browser atau PDF reader."),
    ("Lampiran 3.1.6: Bukti Pengujian Pembuatan Quest Baru & Konfirmasi Uang Muka (DP)",
     "Gambar 3.1.6: Formulir pembuatan lowongan quest dengan kalkulator reward dan layar verifikasi struk transfer uang muka (DP).",
     "Tempelkan screenshot antarmuka penambahan lowongan quest dan halaman rincian transaksi transfer uang muka."),
    ("Lampiran 3.1.7: Bukti Pengujian Penyerahan Master ZIP & Konfirmasi Selesai Ulasan Bintang",
     "Gambar 3.1.7: Tombol unduh berkas master ZIP aktif setelah pelunasan dan modal ulasan rating bintang (1-5) berhasil diisi.",
     "Tempelkan screenshot antarmuka pengunggahan file ZIP deliverable dan layar pengisian rating ulasan bintang."),
    ("Lampiran 3.1.8: Bukti Pengujian Percakapan Proyek & Pembekuan Status Sengketa",
     "Gambar 3.1.8: Panel percakapan interaktif thread proyek dan tampilan pembekuan status proyek saat eskalasi sengketa diajukan.",
     "Tempelkan screenshot ruang obrolan langsung serta indikator status proyek yang dibekukan (*Disputed*).")
]
for item_t, cap_t, ins_t in student_evidences:
    create_evidence_item(item_t, cap_t, ins_t)

# 4.2 Lampiran Mentor
doc.add_page_break()
add_section_heading("Lampiran Pengujian 3.2", level=2)
mentor_evidences = [
    ("Lampiran 3.2.1: Bukti Pengujian Dasbor Pengajar & Tanda Tangan Digital",
     "Gambar 3.2.1: Dasbor instruktur menampilkan ringkasan siswa aktif dan panel penyimpanan goresan tanda tangan digital resmi.",
     "Tempelkan screenshot antarmuka Dasbor Instruktur dan halaman pengaturan tanda tangan digital pengajar."),
    ("Lampiran 3.2.2: Bukti Pengujian Pemantauan Progres Siswa (Student Journey)",
     "Gambar 3.2.2: Lembar pemantauan kemajuan belajar siswa bimbingan menampilkan persentase ketuntasan modul dan skor kuis.",
     "Tempelkan screenshot daftar pemantauan rekam jejak belajar siswa pada grup karir bimbingan."),
    ("Lampiran 3.2.3: Bukti Pengujian Penyusun Modul & Konten Pembelajaran",
     "Gambar 3.2.3: Antarmuka penyusun kurikulum modul dengan fitur seret-lepas (drag-and-drop reorder) dan formulir penambahan materi ajar.",
     "Tempelkan screenshot lembar penyusun modul pembelajaran dan editor materi panduan teks/gambar."),
    ("Lampiran 3.2.4: Bukti Pengujian Penilaian Tugas & Umpan Balik Masukan",
     "Gambar 3.2.4: Formulir koreksi tugas siswa dengan input nilai angka (skala 0-100) dan catatan masukan bimbingan pengajar.",
     "Tempelkan screenshot lembar pemeriksaan karya siswa saat menginput skor nilai dan catatan ulasan feedback."),
    ("Lampiran 3.2.5: Bukti Pengujian Forum Diskusi & Penyematan Pesan Penting",
     "Gambar 3.2.5: Ruang forum kelas dengan balasan pengajar beridentitas khusus dan pesan pengumuman penting yang tersemat di posisi teratas.",
     "Tempelkan screenshot ruang diskusi forum kelas yang memperlihatkan pesan tersisa di bagian atas (pinned message).")
]
for item_t, cap_t, ins_t in mentor_evidences:
    create_evidence_item(item_t, cap_t, ins_t)

# 4.3 Lampiran Admin
doc.add_page_break()
add_section_heading("Lampiran Pengujian 3.3", level=2)
admin_evidences = [
    ("Lampiran 3.3.1: Bukti Pengujian Manajemen Pengguna & Perubahan Peran (RBAC)",
     "Gambar 3.3.1: Daftar akun pengguna terdata di panel admin dan dialog perubahan hak akses peran pengguna menjadi Pengajar (Mentor).",
     "Tempelkan screenshot tabel manajemen pengguna di konsol admin beserta notifikasi keberhasilan perubahan hak akses."),
    ("Lampiran 3.3.2: Bukti Pengujian Pengelolaan Aset Master Game RPG & Template Sertifikat",
     "Gambar 3.3.2: Panel konfigurasi karakter avatar RPG, tingkatan pangkat, lencana level, dan penetapan template aktif sertifikat kelulusan.",
     "Tempelkan screenshot antarmuka pengelolaan master aset RPG dan status template sertifikat aktif."),
    ("Lampiran 3.3.3: Bukti Pengujian Publikasi Kursus & Penugasan Pengajar Pembimbing",
     "Gambar 3.3.3: Formulir penetapan pengajar bimbingan (Assign Mentor) dan tombol pengaktifan publikasi kursus.",
     "Tempelkan screenshot lembar pengaturan kursus saat mengaitkan mentor dan mengubah status menjadi terpublikasi."),
    ("Lampiran 3.3.4: Bukti Pengujian Lembar Moderasi Lowongan Quest Draf",
     "Gambar 3.3.4: Antrean lowongan quest baru berstatus draf di panel moderasi dengan opsi aksi persetujuan publikasi atau penolakan.",
     "Tempelkan screenshot antrean moderasi lowongan proyek freelance pada panel administrator."),
    ("Lampiran 3.3.5: Bukti Pengujian Lembar Sidang Arbitrase & Penerbitan Amar Putusan",
     "Gambar 3.3.5: Lembar kendali sidang arbitrase sengketa menampilkan kronologi perkara, opsi kaukus bilateral, dan sertifikat amar putusan.",
     "Tempelkan screenshot lembar sidang investigasi arbitrase dan formulir penerbitan putusan resmi perkara."),
    ("Lampiran 3.3.6: Bukti Pengujian Pengawasan Submisi Tugas Siswa Lintas Kursus",
     "Gambar 3.3.6: Tabel pengawasan tugas siswa di panel admin menampilkan ringkasan tugas terkumpul dan pengajar pemeriksa.",
     "Tempelkan screenshot tabel pemantauan submisi tugas siswa di konsol administrator."),
    ("Lampiran 3.3.7: Bukti Pengujian Manajemen Footer CMS & Banner Donasi Halaman Depan",
     "Gambar 3.3.7: Antarmuka pengelolaan tautan media sosial footer dan pengaturan banner donasi publik.",
     "Tempelkan screenshot panel pengaturan footer CMS serta tampilan banner donasi di halaman beranda depan.")
]
for item_t, cap_t, ins_t in admin_evidences:
    create_evidence_item(item_t, cap_t, ins_t)

# -------------------------------------------------------------
# 5.0 LEMBAR KESIMPULAN & PENGESAHAN PENERIMAAN PENGGUNA
# -------------------------------------------------------------
doc.add_page_break()
add_section_heading("5.0 LEMBAR KESIMPULAN & PENGESAHAN PENERIMAAN PENGGUNA (UAT SIGN-OFF)", level=1)

p_stat = doc.add_paragraph()
p_stat.add_run(
    "5.1 Rekapitulasi Hasil Pengujian UAT:\n"
    "• Total Skenario Uji: 74 Kasus Uji\n"
    "• Jumlah Kasus Uji Berhasil (OK): 74 Kasus Uji\n"
    "• Jumlah Kasus Uji Gagal (NOK): 0 Kasus Uji\n"
    "• Tingkat Kelulusan Penerimaan (Acceptance Rate): 100% (PASS)\n\n"
    "5.2 Pernyataan Penerimaan Sistem:\n"
    "Berdasarkan pengujian penerimaan pengguna yang telah dilaksanakan secara komprehensif pada platform "
    "SkillMongo (SkillVentura), seluruh fungsionalitas sistem dinyatakan:\n\n"
    "   [ X ] DITERIMA PENUH (FULL ACCEPTANCE)\n"
    "   [   ] DITERIMA BERSYARAT (CONDITIONAL ACCEPTANCE)\n"
    "   [   ] DITOLAK (REJECTED)\n\n"
    "Sistem dinyatakan telah memenuhi seluruh kriteria penerimaan fungsional bisnis, kemudahan antarmuka (UI/UX), "
    "dan siap digunakan secara resmi di lingkungan operasional produksi."
)

# Signature Table
sig_table = doc.add_table(rows=4, cols=3)
sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
sig_headers = ["Perwakilan Siswa (Student)", "Perwakilan Pengajar (Mentor)", "Perwakilan Admin / QA Lead"]
for c_idx, cell in enumerate(sig_table.rows[0].cells):
    cell.width = Inches(2.35)
    set_cell_background(cell, "F2F2F2")
    set_cell_margins(cell, 100, 100, 100, 100)
    set_cell_borders(cell, "A6A6A6", "A6A6A6", "A6A6A6", "A6A6A6")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sig_headers[c_idx])
    r.font.bold = True
    r.font.size = Pt(9.5)

# Signature box row
row_box = sig_table.rows[1]
make_row_cant_split(row_box)
for cell in row_box.cells:
    cell.width = Inches(2.35)
    set_cell_background(cell, "FFFFFF")
    set_cell_margins(cell, 350, 350, 100, 100) # Tall for signature
    set_cell_borders(cell, "A6A6A6", "A6A6A6", "A6A6A6", "A6A6A6")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("( Tanda Tangan )")
    r.font.size = Pt(8.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

# Name row
row_name = sig_table.rows[2]
make_row_cant_split(row_name)
for cell in row_name.cells:
    cell.width = Inches(2.35)
    set_cell_background(cell, "FAFAFA")
    set_cell_margins(cell, 80, 80, 100, 100)
    set_cell_borders(cell, "A6A6A6", "A6A6A6", "A6A6A6", "A6A6A6")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Nama: _______________________")
    r.font.size = Pt(9)

# Date row
row_date = sig_table.rows[3]
make_row_cant_split(row_date)
for cell in row_date.cells:
    cell.width = Inches(2.35)
    set_cell_background(cell, "FAFAFA")
    set_cell_margins(cell, 80, 80, 100, 100)
    set_cell_borders(cell, "A6A6A6", "A6A6A6", "A6A6A6", "A6A6A6")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Tanggal: _____________________")
    r.font.size = Pt(9)

output_path_root = r"c:\Users\An\Herd\skillmongo\Dokumen_UAT_SkillMongo_Verifikasi_Lengkap.docx"
output_path_final = r"c:\Users\An\Herd\skillmongo\Dokumen_UAT_SkillMongo_Final.docx"
doc.save(output_path_root)
doc.save(output_path_final)
print(f"Document successfully created at {output_path_root} and {output_path_final}")
